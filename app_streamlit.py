"""
ATS Intelligence Engine - Sandbox Demo
"""

import streamlit as st
import tempfile
import os
import shutil
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

st.set_page_config(
    page_title="ATS Intelligence",
    page_icon="⚡",
    layout="wide",
    initial_sidebar_state="collapsed",
)

st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@300;400;500;600;700&display=swap');

html, body, [class*="css"] { font-family: 'Inter', sans-serif; }
#MainMenu, footer, header { visibility: hidden; }
.block-container { padding-top: 2rem; padding-bottom: 2rem; max-width: 900px; }

.hero {
    background: linear-gradient(135deg, #0f0f23 0%, #1a1a3e 50%, #0d1b2a 100%);
    border-radius: 16px;
    padding: 2.5rem 2rem;
    margin-bottom: 2rem;
    text-align: center;
    border: 1px solid rgba(99,102,241,0.3);
}
.card-label {
    font-size: 0.72rem;
    font-weight: 600;
    letter-spacing: 0.08em;
    text-transform: uppercase;
    color: #64748b;
    margin-bottom: 0.5rem;
}
.stats-row {
    display: flex;
    gap: 0.8rem;
    flex-wrap: wrap;
    margin: 1.2rem 0;
}
.stat-pill {
    background: #f8fafc;
    border: 1px solid #e2e8f0;
    border-radius: 8px;
    padding: 0.6rem 1rem;
    flex: 1;
    min-width: 120px;
    text-align: center;
}
.stat-pill .stat-val { font-size: 1.5rem; font-weight: 700; color: #1e293b; }
.stat-pill .stat-lbl { font-size: 0.72rem; color: #64748b; margin-top: 0.1rem; }

.cand-card {
    background: #f8fafc;
    border: 1px solid #e2e8f0;
    border-left: 4px solid #6366f1;
    border-radius: 8px;
    padding: 1rem 1.2rem;
    margin-bottom: 0.7rem;
}
.cand-header {
    display: flex;
    justify-content: space-between;
    align-items: center;
    margin-bottom: 0.4rem;
}
.cand-rank { font-size: 0.72rem; font-weight: 700; color: #6366f1; letter-spacing: 0.06em; text-transform: uppercase; }
.cand-score { background: #eef2ff; color: #4f46e5; font-size: 0.8rem; font-weight: 600; padding: 0.2rem 0.7rem; border-radius: 99px; }
.cand-id { font-size: 0.95rem; font-weight: 600; color: #1e293b; }
.cand-reason { font-size: 0.85rem; color: #475569; margin-top: 0.4rem; line-height: 1.5; }

.score-bar-wrap { margin: 0.3rem 0; }
.score-bar-label { display: flex; justify-content: space-between; font-size: 0.75rem; color: #64748b; margin-bottom: 0.15rem; }
.score-bar-bg { background: #e2e8f0; border-radius: 99px; height: 5px; overflow: hidden; }
.score-bar-fill { height: 100%; background: linear-gradient(90deg, #6366f1, #818cf8); border-radius: 99px; }

.stButton > button {
    background: linear-gradient(135deg, #6366f1, #4f46e5) !important;
    color: white !important;
    border: none !important;
    border-radius: 8px !important;
    font-weight: 600 !important;
    font-size: 0.95rem !important;
    padding: 0.65rem 1.5rem !important;
    width: 100%;
}
.stButton > button:hover { opacity: 0.88 !important; }
.stDownloadButton > button {
    background: #f0fdf4 !important;
    color: #16a34a !important;
    border: 1px solid #bbf7d0 !important;
    border-radius: 8px !important;
    font-weight: 600 !important;
    width: 100%;
}
.divider { border: none; border-top: 1px solid #e2e8f0; margin: 1.5rem 0; }
</style>
""", unsafe_allow_html=True)

# ── Hero ──────────────────────────────────────────────────────────────────────
st.markdown("""
<div class="hero">
    <div style="font-size:2.2rem;font-weight:700;color:#ffffff;letter-spacing:-0.02em;margin-bottom:0.5rem;">
        ATS <span style="color:#818cf8;">Intelligence</span> Engine
    </div>
    <div style="color:#94a3b8;font-size:1rem;">
        Resume ranking across 10 dimensions &nbsp;&middot;&nbsp; Honeypot detection &nbsp;&middot;&nbsp; Behavioral signals
    </div>
</div>
""", unsafe_allow_html=True)

# ── Inputs ────────────────────────────────────────────────────────────────────
col1, col2 = st.columns([1, 1], gap="large")

with col1:
    st.markdown('<div class="card-label">Candidates File</div>', unsafe_allow_html=True)
    uploaded_file = st.file_uploader(
        "Upload candidates file",
        type=["jsonl", "json"],
        label_visibility="collapsed",
    )
    st.caption("JSON or JSONL · max 100,000 candidates")
    st.markdown('<div class="card-label" style="margin-top:0.8rem;">Top Candidates to Rank</div>', unsafe_allow_html=True)
    top_k = st.number_input("Top candidates to rank", min_value=10, max_value=100, value=50, step=1, label_visibility="collapsed")

with col2:
    st.markdown('<div class="card-label">Job Description</div>', unsafe_allow_html=True)
    jd_file = st.file_uploader(
        "Upload JD file",
        type=["txt", "docx", "pdf"],
        label_visibility="collapsed",
        key="jd_upload",
    )
    st.caption("Upload .txt / .docx / .pdf  —  or type below")
    jd_text = st.text_area(
        "Job Description",
        height=110,
        placeholder="Paste the job description here…",
        label_visibility="collapsed",
    )

st.markdown("<hr class='divider'>", unsafe_allow_html=True)
run_btn = st.button("⚡  Run Ranking", use_container_width=True)

# ── Pipeline ──────────────────────────────────────────────────────────────────
if run_btn:
    import time
    wall_start = time.time()

    if not uploaded_file:
        st.error("Upload a candidates file first.")
        st.stop()

    # Resolve JD
    final_jd = jd_text.strip()
    if jd_file is not None:
        fname = jd_file.name.lower()
        if fname.endswith(".docx"):
            try:
                import docx
                from io import BytesIO
                doc = docx.Document(BytesIO(jd_file.getvalue()))
                final_jd = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
            except ImportError:
                st.error("python-docx not installed. Run: pip install python-docx")
                st.stop()
        elif fname.endswith(".pdf"):
            try:
                import pdfplumber
                from io import BytesIO
                with pdfplumber.open(BytesIO(jd_file.getvalue())) as pdf:
                    final_jd = "\n".join(page.extract_text() or "" for page in pdf.pages)
            except ImportError:
                st.error("pdfplumber not installed. Run: pip install pdfplumber")
                st.stop()
        else:
            final_jd = jd_file.getvalue().decode("utf-8", errors="ignore")

    if not final_jd:
        st.error("Enter a job description or upload a JD file.")
        st.stop()

    progress = st.progress(0)
    status   = st.empty()

    try:
        from generate_submission import load_candidates_jsonl, prefilter_candidates_tfidf, save_temp_resumes
        from app.services.batch.pipeline import BatchPipeline
        from app.services.reasoning_generator import ReasoningGenerator, parse_jd_requirements
        from app.services.behavioral_signals.scorer import apply_behavioral_multiplier
        from app.services.honeypot_detector import HoneypotDetector

        # 1 — load
        status.info("📂  Loading candidates…")
        progress.progress(8)
        import json as _json
        raw_text = uploaded_file.getvalue().decode("utf-8", errors="ignore").strip()
        jsonl_tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".jsonl", mode="w", encoding="utf-8")
        if raw_text.startswith("["):
            for item in _json.loads(raw_text):
                jsonl_tmp.write(_json.dumps(item) + "\n")
        else:
            for line in raw_text.splitlines():
                if line.strip():
                    jsonl_tmp.write(line.strip() + "\n")
        jsonl_tmp.close()
        candidates = load_candidates_jsonl(jsonl_tmp.name)
        os.unlink(jsonl_tmp.name)

        if not candidates:
            st.error("No valid candidates found in the file.")
            st.stop()
        if len(candidates) > 100000:
            candidates = candidates[:100000]
            st.warning("Truncated to first 100,000 candidates.")

        st.success(f"✅ Loaded **{len(candidates)}** candidates")

        # 2 — honeypots
        status.info("🍯  Detecting honeypots…")
        progress.progress(20)
        detector = HoneypotDetector(strict_mode=False)
        valid_data, honeypots = detector.filter_honeypots([c["candidate_data"] for c in candidates])
        valid_ids  = {c.get("candidate_id") for c in valid_data}
        candidates = [c for c in candidates if c["candidate_id"] in valid_ids]

        PRE_FILTER_K = 2000
        if len(candidates) > PRE_FILTER_K:
            status.info(f"Fast shortlisting {len(candidates):,} candidates...")
            progress.progress(28)
            candidates = prefilter_candidates_tfidf(candidates, final_jd, PRE_FILTER_K)

        # 3 — pipeline
        status.info("⚙️  Running screening pipeline…")
        progress.progress(35)
        temp_dir = tempfile.mkdtemp()
        save_temp_resumes(candidates, temp_dir)
        pipeline = BatchPipeline()
        result   = pipeline.run_tiered(
            resume_dir=temp_dir,
            job_description=final_jd,
            top_k=min(top_k, len(candidates)),
        )

        # Build lookup
        candidate_lookup = {c["candidate_id"]: c["candidate_data"] for c in candidates}

        # 4 — behavioral signals
        status.info("📊  Applying behavioral signals…")
        progress.progress(60)
        jd_requirements = parse_jd_requirements(final_jd)

        for rc in result["candidates"]:
            cid = rc["resume_id"].replace(".txt", "")
            full_data = candidate_lookup.get(cid, {})
            beh = apply_behavioral_multiplier(
                base_score=rc["overall_score"],
                redrob_signals=full_data.get("redrob_signals", {}),
                jd_requirements=jd_requirements,
            )
            rc["overall_score"]         = beh["final_score"]
            rc["behavioral_multiplier"] = beh["behavioral_multiplier"]

        result["candidates"].sort(key=lambda x: x["overall_score"], reverse=True)
        for i, rc in enumerate(result["candidates"], 1):
            rc["rank"] = i
            rc["resume_id"] = rc["resume_id"].replace(".txt", "")

        # 5 — reasoning
        status.info("💬  Generating reasoning…")
        progress.progress(80)
        reasoning_gen = ReasoningGenerator(jd_requirements)
        for rc in result["candidates"]:
            full_data = candidate_lookup.get(rc["resume_id"], {})
            rc["reasoning"]      = reasoning_gen.generate(
                candidate=full_data, rank=rc["rank"], scores=rc["all_scores"]
            )
            rc["candidate_data"] = full_data

        # 6 — build CSV (spec columns: candidate_id, rank, score, reasoning)
        progress.progress(95)
        csv_lines = ["candidate_id,rank,score,reasoning"]
        for rc in result["candidates"]:
            cid       = rc.get("resume_id", "")
            score     = rc.get("overall_score", 0) / 100
            reasoning = rc.get("reasoning", "").replace('"', "'").replace("\n", " ")
            csv_lines.append(f'{cid},{rc["rank"]},{score:.4f},"{reasoning}"')
        csv_data = "\n".join(csv_lines)

        shutil.rmtree(temp_dir, ignore_errors=True)
        progress.progress(100)
        status.empty()

        # ── Stats ────────────────────────────────────────────────────────────
        wall_elapsed = time.time() - wall_start
        full_data_lookup = {c["candidate_id"]: c["candidate_data"] for c in candidates}

        st.markdown(f"""
        <div class="stats-row">
            <div class="stat-pill"><div class="stat-val">{result['total_resumes_processed']}</div><div class="stat-lbl">Screened</div></div>
            <div class="stat-pill"><div class="stat-val">{len(result['candidates'])}</div><div class="stat-lbl">Ranked</div></div>
            <div class="stat-pill"><div class="stat-val">{len(honeypots)}</div><div class="stat-lbl">Honeypots removed</div></div>
            <div class="stat-pill"><div class="stat-val">{wall_elapsed/60:.1f}m</div><div class="stat-lbl">Total time</div></div>
        </div>
        """, unsafe_allow_html=True)

        # ── Excel builder (spec columns: candidate_id, rank, score, reasoning) ──
        def build_excel_reasoning(rc, full_data):
            profile  = full_data.get("profile", {})
            skills   = full_data.get("skills", [])
            yoe      = profile.get("years_of_experience", 0)
            title    = profile.get("current_title", "")
            company  = profile.get("current_company", "")
            industry = profile.get("current_industry", "")
            scores   = rc.get("all_scores", {})
            rank     = rc.get("rank", 0)
            overall  = rc.get("overall_score", 0)
            sem      = scores.get("semantic_fit", 0)
            cgrow    = scores.get("career_growth", 0)

            prof_order = {"expert": 4, "advanced": 3, "intermediate": 2, "beginner": 1}
            top_skills = sorted(skills, key=lambda x: prof_order.get(x.get("proficiency", ""), 0), reverse=True)
            adv = [s["name"] for s in top_skills if s.get("proficiency") in ("expert", "advanced")][:3]
            mid = [s["name"] for s in top_skills if s.get("proficiency") == "intermediate"][:2]
            skill_str = ", ".join(adv) if adv else ", ".join(mid) if mid else "general skills"

            line1 = f"{yoe:.0f} yrs exp as {title} at {company} ({industry}); strong in {skill_str}."

            concerns  = []
            strengths = []
            if sem > 75:    strengths.append("strong JD alignment")
            elif sem > 50:  strengths.append("moderate JD fit")
            else:           concerns.append("limited JD alignment")
            if cgrow > 65:  strengths.append("clear career growth")
            elif cgrow < 40: concerns.append("flat career trajectory")
            if yoe < 3:     concerns.append(f"low experience ({yoe:.0f} yrs, JD wants 5-9)")
            elif yoe > 12:  strengths.append("deep domain experience")

            tone = "Strong fit" if rank <= 15 else ("Good fit" if rank <= 40 else ("Moderate fit" if rank <= 70 else "Weak fit"))
            concern_str  = f" Concerns: {'; '.join(concerns)}." if concerns else ""
            strength_str = f" {'; '.join(strengths).capitalize()}." if strengths else ""
            line2 = f"{tone} — overall score {overall:.1f}/100.{strength_str}{concern_str}"

            return f"{line1}\n{line2}"

        import io
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Rankings"
            headers = ["candidate_id", "rank", "score", "reasoning"]
            hfill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
            for col, h in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=h)
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = hfill
                cell.alignment = Alignment(horizontal="center")
            for rc in result["candidates"]:
                cid       = rc.get("resume_id", "")
                ws.append([
                    cid,
                    rc.get("rank"),
                    round(rc.get("overall_score", 0) / 100, 4),
                    rc.get("reasoning", ""),
                ])
            ws.column_dimensions["A"].width = 18
            ws.column_dimensions["B"].width = 8
            ws.column_dimensions["C"].width = 10
            ws.column_dimensions["D"].width = 90
            for row in ws.iter_rows(min_row=2):
                row[3].alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[1].height = 20
            excel_buf = io.BytesIO()
            wb.save(excel_buf)
            excel_data = excel_buf.getvalue()
            excel_ok = True
        except ImportError:
            excel_ok = False

        # ── Downloads ─────────────────────────────────────────────────────────
        col_csv, col_xl = st.columns(2)
        with col_csv:
            st.download_button("📥  Download CSV", data=csv_data,
                file_name="ranking_results.csv", mime="text/csv",
                use_container_width=True)
        with col_xl:
            if excel_ok:
                st.download_button("📊  Download Excel", data=excel_data,
                    file_name="ranking_results.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True)

        # ── Candidate cards ───────────────────────────────────────────────────
        st.markdown("<hr class='divider'>", unsafe_allow_html=True)
        st.markdown(f"#### 🏆 Top {len(result['candidates'])} Candidates")

        SKIP_DIMS = {"resilience", "narrative_coherence"}
        for rc in result["candidates"]:
            cid       = rc.get("resume_id", "")
            full_data = full_data_lookup.get(cid, {})
            name      = full_data.get("profile", {}).get("anonymized_name", "")
            scores    = rc.get("all_scores", {})

            bars_html = ""
            filtered  = [(d, v) for d, v in scores.items() if d.lower() not in SKIP_DIMS]
            for dim, val in filtered[:5]:
                pct = min(max(float(val), 0), 100)
                bars_html += (
                    f'<div class="score-bar-wrap">'
                    f'<div class="score-bar-label"><span>{dim.replace("_"," ").title()}</span><span>{val:.0f}</span></div>'
                    f'<div class="score-bar-bg"><div class="score-bar-fill" style="width:{pct:.1f}%"></div></div>'
                    f'</div>'
                )

            card_html = (
                f'<div class="cand-card">'
                f'<div class="cand-header">'
                f'<span class="cand-rank">Rank #{rc["rank"]}</span>'
                f'<span class="cand-score">Score {rc["overall_score"]:.1f}</span>'
                f'</div>'
                f'<div class="cand-id">{cid} — {name}</div>'
                f'<div class="cand-reason">{rc.get("reasoning", "—")}</div>'
                + bars_html +
                f'</div>'
            )
            st.markdown(card_html, unsafe_allow_html=True)

    except Exception as exc:
        import traceback
        status.empty()
        st.error(f"Something went wrong: {exc}")
        with st.expander("Error details"):
            st.code(traceback.format_exc())
        try:
            shutil.rmtree(temp_dir, ignore_errors=True)
        except Exception:
            pass
