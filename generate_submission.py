#!/usr/bin/env python3
"""
Hackathon Submission Generator

Usage:
    python generate_submission.py \
        --candidates ./docs/candidates.jsonl \
        --jd ./docs/job_description.txt \
        --output team_xxx.csv
"""

import argparse
import csv
import json
import sys
import os
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from app.services.batch.pipeline import BatchPipeline
from app.services.reasoning_generator import ReasoningGenerator, parse_jd_requirements
from app.services.csv_writer import write_submission_csv
from app.services.behavioral_signals.scorer import apply_behavioral_multiplier
from app.services.honeypot_detector import HoneypotDetector


def load_candidates_jsonl(jsonl_path: str) -> list[dict]:
    candidates = []
    with open(jsonl_path, 'r', encoding='utf-8') as f:
        for line_num, line in enumerate(f, 1):
            line = line.strip()
            if not line:
                continue
            try:
                data = json.loads(line)
                resume_text = _candidate_to_text(data)
                candidates.append({
                    'candidate_id': data.get('candidate_id', f'CAND_{line_num:07d}'),
                    'resume_text': resume_text,
                    'candidate_data': data
                })
            except json.JSONDecodeError as e:
                print(f"Warning: Skipping invalid JSON on line {line_num}: {e}")
    return candidates


def _candidate_to_text(candidate: dict) -> str:
    parts = []
    if 'profile' in candidate:
        profile = candidate['profile']
        parts.append(profile.get('headline', ''))
        parts.append(profile.get('summary', ''))
        parts.append(profile.get('current_title', ''))
        parts.append(profile.get('current_company', ''))
        parts.append(profile.get('location', ''))
    if 'career_history' in candidate:
        for role in candidate['career_history']:
            parts.append(role.get('title', ''))
            parts.append(role.get('company', ''))
            parts.append(role.get('description', ''))
    if 'skills' in candidate:
        skill_names = [s.get('name', '') for s in candidate['skills']]
        parts.append(' '.join(skill_names))
    if 'education' in candidate:
        for edu in candidate['education']:
            parts.append(edu.get('institution', ''))
            parts.append(edu.get('degree', ''))
            parts.append(edu.get('field_of_study', ''))
    return ' '.join(filter(None, parts))


def save_temp_resumes(candidates: list[dict], temp_dir: str) -> str:
    """Write shortlisted resumes as one JSONL file for fast batch loading.

    The batch pipeline already supports a `resumes.jsonl` file. Using that
    avoids creating thousands of tiny .txt files, which is especially slow on
    Windows and was the main bottleneck for 100k-candidate runs.
    """
    temp_path = Path(temp_dir)
    temp_path.mkdir(parents=True, exist_ok=True)
    jsonl_path = temp_path / "resumes.jsonl"
    with open(jsonl_path, 'w', encoding='utf-8') as f:
        for candidate in candidates:
            f.write(json.dumps({
                'id': candidate['candidate_id'],
                'resume_text': candidate['resume_text'],
            }, ensure_ascii=False) + '\n')
    return str(temp_path)


def save_temp_resumes_as_txt(candidates: list[dict], temp_dir: str) -> str:
    """Legacy writer kept for debugging or comparison."""
    temp_path = Path(temp_dir)
    temp_path.mkdir(parents=True, exist_ok=True)
    for candidate in candidates:
        cid = candidate['candidate_id']
        file_path = temp_path / f"{cid}.txt"
        with open(file_path, 'w', encoding='utf-8') as f:
            f.write(candidate['resume_text'])
    return str(temp_path)


def prefilter_candidates_tfidf(
    candidates: list[dict],
    jd_text: str,
    limit: int = 2000,
    max_features: int = 8000,
) -> list[dict]:
    """Fast in-memory shortlist before the heavier ranking pipeline."""
    if len(candidates) <= limit:
        return candidates

    from sklearn.feature_extraction.text import TfidfVectorizer
    from sklearn.metrics.pairwise import cosine_similarity
    import numpy as np

    texts = [c['resume_text'] for c in candidates]
    vec = TfidfVectorizer(max_features=max_features, stop_words='english')
    mat = vec.fit_transform(texts)
    jd_vec = vec.transform([jd_text])
    sims = cosine_similarity(jd_vec, mat).flatten()
    top_idx = np.argsort(sims)[::-1][:limit]
    return [candidates[i] for i in top_idx]


def write_submission_xlsx(candidates: list[dict], output_path: str) -> dict:
    """Write an Excel copy of the submission with the same four spec columns."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError as exc:
        raise RuntimeError("openpyxl is required for Excel output") from exc

    wb = Workbook(write_only=False)
    ws = wb.active
    ws.title = "Rankings"

    headers = ['candidate_id', 'rank', 'score', 'reasoning']
    header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for candidate in sorted(candidates, key=lambda x: x.get('rank', 999)):
        candidate_id = candidate.get('resume_id') or candidate.get('candidate_id')
        score = candidate.get('overall_score', 0) / 100
        reasoning = candidate.get('reasoning', '').replace('\r', '').strip()
        ws.append([candidate_id, candidate.get('rank'), round(score, 4), reasoning])

    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 8
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 90
    for row in ws.iter_rows(min_row=2):
        row[3].alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(output_path)
    return {'success': True, 'output_path': output_path, 'rows_written': len(candidates)}


def build_reasoning(rc: dict, full_data: dict, scores: dict) -> str:
    """Compatibility wrapper; actual logic lives in ReasoningGenerator."""
    return ReasoningGenerator().generate(
        candidate=full_data,
        rank=rc.get('rank', 0),
        scores=scores | {'overall_score': rc.get('overall_score', 0)},
    )


def main():
    parser = argparse.ArgumentParser(description="Generate hackathon submission CSV")
    parser.add_argument('--candidates', required=True)
    parser.add_argument('--jd', '--jd-file', dest='jd_file', required=True)
    parser.add_argument('--output', required=True)
    parser.add_argument('--excel-output', default=None)
    parser.add_argument('--top-k', type=int, default=100)
    parser.add_argument('--tiered', action='store_true', default=True)
    parser.add_argument('--workers', type=int, default=8)
    parser.add_argument('--temp-dir', default='./temp_resumes')
    parser.add_argument('--validate-only', action='store_true')
    parser.add_argument('--verbose', action='store_true')
    args = parser.parse_args()

    if args.top_k != 100:
        args.top_k = 100

    print(f"{'='*80}")
    print("HACKATHON SUBMISSION GENERATOR")
    print(f"{'='*80}\n")

    # Step 1: Load JD
    print(f"[1/7] Loading job description from {args.jd_file}...")
    with open(args.jd_file, 'r', encoding='utf-8') as f:
        jd_text = f.read()
    jd_requirements = parse_jd_requirements(jd_text)
    print(f"      Skills: {len(jd_requirements.get('skills', []))}, "
          f"Exp: {jd_requirements.get('min_experience', 0)}-{jd_requirements.get('max_experience', 0)} yrs")

    # Step 2: Load candidates
    print(f"\n[2/7] Loading candidates from {args.candidates}...")
    candidates = load_candidates_jsonl(args.candidates)
    print(f"      Loaded {len(candidates)} candidates")

    # Step 2.5: Honeypot filter
    print(f"\n[2.5/7] Detecting honeypots...")
    honeypot_detector = HoneypotDetector(strict_mode=False)
    valid_candidates, honeypots = honeypot_detector.filter_honeypots(
        [c['candidate_data'] for c in candidates]
    )
    print(f"      Removed {len(honeypots)} honeypots, kept {len(valid_candidates)} valid")
    valid_ids = {c.get('candidate_id') for c in valid_candidates}
    candidates = [c for c in candidates if c['candidate_id'] in valid_ids]

    # Step 2.8: TF-IDF pre-filter to top 2000
    PRE_FILTER_K = 2000
    if len(candidates) > PRE_FILTER_K:
        print(f"\n[2.8/7] TF-IDF pre-filtering {len(candidates)} -> {PRE_FILTER_K}...")
        candidates = prefilter_candidates_tfidf(candidates, jd_text, PRE_FILTER_K)
        print(f"      Done - {len(candidates)} candidates going to pipeline")

    # Build lookup before pipeline
    candidate_lookup = {c['candidate_id']: c['candidate_data'] for c in candidates}

    # Step 3: Save temp files
    print(f"\n[3/7] Saving {len(candidates)} resumes to one JSONL batch file...")
    temp_dir = save_temp_resumes(candidates, args.temp_dir)

    # Step 4: Run pipeline
    print(f"\n[4/7] Running screening pipeline...")
    pipeline = BatchPipeline()
    result = pipeline.run_tiered(
        resume_dir=temp_dir,
        job_description=jd_text,
        top_k=args.top_k,
        num_workers=args.workers
    )
    print(f"      OK {result['total_resumes_processed']} resumes in {result['elapsed_seconds']:.1f}s")
    print(f"      OK Top {len(result['candidates'])} ranked")

    # Step 4.5: Behavioral signals
    print(f"\n[4.5/7] Applying behavioral signals...")
    for rc in result['candidates']:
        cid = rc['resume_id'].replace('.txt', '')
        full_data = candidate_lookup.get(cid, {})
        beh = apply_behavioral_multiplier(
            base_score=rc['overall_score'],
            redrob_signals=full_data.get('redrob_signals', {}),
            jd_requirements=jd_requirements
        )
        rc['base_score']            = rc['overall_score']
        rc['overall_score']         = beh['final_score']
        rc['behavioral_multiplier'] = beh['behavioral_multiplier']

    result['candidates'].sort(key=lambda x: x['overall_score'], reverse=True)
    for i, rc in enumerate(result['candidates'], 1):
        rc['rank'] = i
    print(f"      OK Done")

    # Step 5: Generate reasoning
    print(f"\n[5/7] Generating reasoning...")
    reasoning_gen = ReasoningGenerator(jd_requirements)
    for rc in result['candidates']:
        cid = rc['resume_id'].replace('.txt', '')
        full_data = candidate_lookup.get(cid, {})
        rc['reasoning']      = reasoning_gen.generate(
            candidate=full_data,
            rank=rc.get('rank', 0),
            scores=rc.get('all_scores', {}) | {'overall_score': rc.get('overall_score', 0)}
        )
        rc['candidate_data'] = full_data
        rc['resume_id']      = cid  # strip .txt permanently

    if args.verbose:
        print(f"\n      Sample (Rank 1): {result['candidates'][0]['reasoning']}")

    # Step 6: Honeypot rate check
    print(f"\n[6/7] Checking honeypot rate in top 100...")
    top_100_data = [candidate_lookup.get(c['resume_id'], {}) for c in result['candidates'][:100]]
    hcheck = honeypot_detector.calculate_honeypot_rate(top_100_data)
    status = "PASSES" if hcheck['passes_threshold'] else "FAILS"
    print(f"      {status} - {hcheck['honeypot_rate']:.1f}% honeypots in top 100")

    # Step 7: Write CSV / Excel
    print(f"\n[7/7] Writing CSV to {args.output}...")
    try:
        write_result = write_submission_csv(
            candidates=result['candidates'],
            output_path=args.output,
            validate_only=args.validate_only
        )
        if not args.validate_only:
            print(f"      OK {write_result['rows_written']} rows written")
            stats = write_result['validation']['stats']
            print(f"      Score range: {stats['min_score']:.4f} - {stats['max_score']:.4f}")
            print(f"      Empty reasoning: {stats['empty_reasoning_count']}")
            if args.excel_output:
                xlsx_result = write_submission_xlsx(result['candidates'], args.excel_output)
                print(f"      OK Excel written: {xlsx_result['output_path']}")
    except Exception as e:
        print(f"      Error: {e}")
        sys.exit(1)

    # Cleanup
    import shutil
    if os.path.exists(temp_dir) and temp_dir.startswith('./temp'):
        shutil.rmtree(temp_dir)

    print(f"\n{'='*80}")
    print("DONE!")
    print(f"{'='*80}")
    print(f"Output: {args.output}")
    print(f"Candidates: {len(result['candidates'])}")
    print(f"Time: {result['elapsed_seconds']:.1f}s")
    print(f"\nReady to submit!")


if __name__ == '__main__':
    main()
